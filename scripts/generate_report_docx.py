"""Generate a readable Word report from structured course-analysis data."""

from __future__ import annotations

import argparse
import json
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

from learning_config import order_bilingual


def _safe_name(value: str) -> str:
    if isinstance(value, dict):
        value = value.get("en", "course-analysis")
    value = re.sub(r"[<>:\"/\\|?*]+", "", value).strip()
    return re.sub(r"\s+", "-", value) or "course-analysis"


def _xml_safe(value: str) -> str:
    """Remove XML-invalid characters and flatten embedded line breaks."""
    cleaned = []
    for char in value:
        if char in "\r\n\t":
            cleaned.append(" ")
        elif ord(char) >= 32:
            cleaned.append(char)
    return "".join(cleaned)


def _contains_cjk(value: str) -> bool:
    return any("\u4e00" <= char <= "\u9fff" for char in value)


def _bilingual(value: object, mode: str = "en-zh") -> tuple[str, str | None]:
    if isinstance(value, dict):
        en = str(value.get("en", ""))
        zh = str(value.get("zh", "")) if value.get("zh") else ""
        ordered = order_bilingual(en, zh, mode)
        return ordered[0], ordered[1] if len(ordered) > 1 else None
    return str(value), None


def _ensure_bilingual_styles(document: Document) -> None:
    """Create the indented continuation style used by bilingual list items."""
    if "Bilingual Chinese" in document.styles:
        return
    style = document.styles.add_style("Bilingual Chinese", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles["Normal"]
    style.paragraph_format.left_indent = Inches(0.3)
    style.paragraph_format.space_after = Pt(4)


def _add_heading_pair(document: Document, english: str, chinese: str | None, level: int, mode: str) -> None:
    first, second = _bilingual({"en": english, "zh": chinese or ""}, mode)
    document.add_heading(_xml_safe(first), level=level)
    if second:
        document.add_heading(_xml_safe(second), level=level)


def _add_bilingual_paragraph(document: Document, value: object, mode: str = "en-zh", style: str | None = None):
    first, second = _bilingual(value, mode)
    paragraph = document.add_paragraph(style=style)
    if second:
        paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(_xml_safe(first))
    run.font.name = "Microsoft YaHei" if _contains_cjk(first) else "Aptos"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), "Microsoft YaHei" if _contains_cjk(first) else "Aptos"
    )
    if second:
        chinese_style = "Bilingual Chinese" if style == "List Bullet" else style
        chinese = document.add_paragraph(style=chinese_style)
        chinese_run = chinese.add_run(_xml_safe(second))
        chinese_run.font.name = "Microsoft YaHei" if _contains_cjk(second) else "Aptos"
        chinese_run._element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "Microsoft YaHei" if _contains_cjk(second) else "Aptos"
        )
    return paragraph


def generate_report(course_root: str | Path, report: dict) -> Path:
    root = Path(course_root).resolve()
    output_dir = root / ".course-assistant" / "reports"
    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    filename = report.get("output_filename") or _safe_name(report.get("title", "analysis"))
    output = output_dir / f"{stamp}-{_safe_name(str(filename))}.docx"

    document = Document()
    language_mode = str(report.get("language_mode", "en-zh"))
    _ensure_bilingual_styles(document)
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    title_en, title_zh = _bilingual(report.get("title", "Course Analysis Report"), language_mode)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run(_xml_safe(title_en))
    if title_zh:
        chinese_title = document.add_paragraph(style="Subtitle")
        chinese_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        chinese_title.add_run(_xml_safe(title_zh))

    course_value = _xml_safe(str(report.get("course", root.name)))
    generated_value = _xml_safe(str(report.get("generated", datetime.now().strftime("%Y-%m-%d"))))
    metadata_lines = order_bilingual(f"Course: {course_value}", f"课程：{course_value}", language_mode)
    generated_lines = order_bilingual(f"Generated: {generated_value}", f"生成日期：{generated_value}", language_mode)
    for line in (*metadata_lines, *generated_lines):
        metadata = document.add_paragraph()
        if line.startswith("Course:") or line.startswith("课程：") or line.startswith("Generated:") or line.startswith("生成日期："):
            label_end = line.find(":") + 1 if ":" in line else line.find("：") + 1
            metadata.add_run(line[:label_end]).bold = True
            metadata.add_run(line[label_end:])
        else:
            metadata.add_run(line)

    status_fields = [
        ("Analysis status", "分析状态", report.get("analysis_status")),
        ("Render status", "渲染状态", report.get("render_status")),
        ("Source count", "来源数量", report.get("source_count")),
    ]
    status_fields = [(label, label_zh, value) for label, label_zh, value in status_fields if value is not None]
    if status_fields:
        for label, label_zh, value in status_fields:
            for line in order_bilingual(f"{label}: {value}", f"{label_zh}：{value}", language_mode):
                status = document.add_paragraph()
                label_end = line.find(":") + 1 if ":" in line else line.find("：") + 1
                status.add_run(line[:label_end]).bold = True
                status.add_run(line[label_end:])

    _add_heading_pair(document, "Core Summary", "核心总结", 1, language_mode)
    _add_bilingual_paragraph(document, report.get("summary", ""), language_mode)

    for section_data in report.get("sections", []):
        heading = section_data.get("heading", "Section")
        _add_heading_pair(document, str(heading), section_data.get("heading_zh"), 1, language_mode)
        for item in section_data.get("items", []):
            _add_bilingual_paragraph(document, item, language_mode, style="List Bullet")

    sources = report.get("sources", [])
    if sources:
        _add_heading_pair(document, "Sources", "来源", 1, language_mode)
        for source in sources:
            document.add_paragraph(str(source), style="List Bullet")

    document.save(output)
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("course_root", type=Path)
    parser.add_argument("report_json", type=Path, help="JSON file containing report data")
    args = parser.parse_args()
    report = json.loads(args.report_json.read_text(encoding="utf-8"))
    print(generate_report(args.course_root, report))


if __name__ == "__main__":
    main()
