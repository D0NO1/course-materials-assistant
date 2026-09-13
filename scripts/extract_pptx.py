"""Extract PowerPoint slide text, tables, notes, and object indicators."""

from __future__ import annotations

from pathlib import Path


def extract_pptx(path: str | Path) -> dict:
    source = Path(path)
    result = {
        "file": source.name,
        "path": str(source),
        "format": "powerpoint",
        "status": "extracted",
        "content": {"slides": []},
        "locators": [],
        "warnings": [],
    }
    try:
        from pptx import Presentation
    except ImportError:
        result["status"] = "inventory-only"
        result["warnings"].append("Install python-pptx for PowerPoint extraction.")
        return result
    try:
        presentation = Presentation(str(source))
        for number, slide in enumerate(presentation.slides, start=1):
            texts = []
            tables = []
            objects = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = shape.text.strip()
                    if text:
                        texts.append(text)
                if getattr(shape, "has_table", False):
                    tables.append([[cell.text for cell in row.cells] for row in shape.table.rows])
                objects.append(shape.shape_type)
            notes = ""
            try:
                notes = slide.notes_slide.notes_text_frame.text.strip()
            except Exception:
                pass
            title = slide.shapes.title.text.strip() if slide.shapes.title else ""
            result["content"]["slides"].append({
                "slide": number,
                "title": title,
                "text": texts,
                "notes": notes,
                "tables": tables,
                "objects": [str(item) for item in objects],
            })
            result["locators"].append({"type": "slide", "number": number, "label": title})
    except Exception as exc:
        result["status"] = "error"
        result["warnings"].append(f"PPTX extraction failed: {exc}")
    return result
