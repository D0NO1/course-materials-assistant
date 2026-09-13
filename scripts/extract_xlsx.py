"""Extract workbook sheet structure, headers, formulas, and hidden-sheet flags."""

from __future__ import annotations

from pathlib import Path


def extract_xlsx(path: str | Path) -> dict:
    source = Path(path)
    result = {
        "file": source.name,
        "path": str(source),
        "format": "excel",
        "status": "extracted",
        "content": {"sheets": []},
        "locators": [],
        "warnings": [],
    }
    try:
        from openpyxl import load_workbook
    except ImportError:
        result["status"] = "inventory-only"
        result["warnings"].append("Install openpyxl for Excel extraction.")
        return result
    try:
        workbook = load_workbook(str(source), read_only=True, data_only=False)
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20), values_only=False))
            headers = [cell.value for cell in rows[0]] if rows else []
            formulas = []
            for row in rows:
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formulas.append({"cell": cell.coordinate, "formula": cell.value})
            result["content"]["sheets"].append({
                "name": sheet.title,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
                "headers": headers,
                "formulas": formulas,
                "hidden": sheet.sheet_state != "visible",
            })
            result["locators"].append({"type": "sheet", "label": sheet.title})
        workbook.close()
    except Exception as exc:
        result["status"] = "error"
        result["warnings"].append(f"XLSX extraction failed: {exc}")
    return result
