"""Extract DOCX paragraphs, headings, and tables."""

from __future__ import annotations

from pathlib import Path


def extract_docx(path: str | Path) -> dict:
    source = Path(path)
    result = {
        "file": source.name,
        "path": str(source),
        "format": "word",
        "status": "extracted",
        "content": {"paragraphs": [], "tables": []},
        "locators": [],
        "warnings": [],
    }
    try:
        from docx import Document
    except ImportError:
        result["status"] = "inventory-only"
        result["warnings"].append("Install python-docx for Word extraction.")
        return result
    try:
        document = Document(str(source))
        for number, paragraph in enumerate(document.paragraphs, start=1):
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name if paragraph.style else ""
            item = {"paragraph": number, "style": style, "text": text}
            result["content"]["paragraphs"].append(item)
            locator_type = "heading" if style.lower().startswith("heading") else "paragraph"
            result["locators"].append({"type": locator_type, "number": number, "label": text[:120]})
        for table_number, table in enumerate(document.tables, start=1):
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            result["content"]["tables"].append({"table": table_number, "rows": rows})
            result["locators"].append({"type": "table", "number": table_number})
    except Exception as exc:
        result["status"] = "error"
        result["warnings"].append(f"DOCX extraction failed: {exc}")
    return result
