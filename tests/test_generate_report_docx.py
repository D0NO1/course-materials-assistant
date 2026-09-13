import sys
import tempfile
import unittest
from zipfile import ZipFile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parents[1] / "scripts"))

from generate_report_docx import generate_report


class GenerateReportDocxTests(unittest.TestCase):
    def test_generate_report_includes_status_metadata_when_provided(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            course = Path(temp_dir) / "course"
            course.mkdir()
            report = {
                "title": "Status Test",
                "course": course.name,
                "generated": "2026-09-13",
                "summary": "Summary",
                "analysis_status": "source-backed",
                "render_status": "not-verified",
                "source_count": 3,
            }

            output = generate_report(course, report)

            with ZipFile(output) as archive:
                document_xml = archive.read("word/document.xml").decode("utf-8")
            self.assertIn("Analysis status", document_xml)
            self.assertIn("source-backed", document_xml)
            self.assertIn("Render status", document_xml)
            self.assertIn("not-verified", document_xml)
            self.assertIn("Source count", document_xml)
            self.assertIn("3", document_xml)

    def test_generate_report_writes_a_readable_bilingual_docx(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            course = Path(temp_dir) / "MIS-413-91 - Systems Analysis"
            course.mkdir()
            report = {
                "title": {"en": "Week 1 Core Analysis", "zh": "第一周核心分析"},
                "course": course.name,
                "generated": "2026-09-09",
                "summary": {
                    "en": "The course teaches how to turn a business problem into a system plan.",
                    "zh": "本课程学习如何把业务问题转化为系统方案。",
                },
                "sections": [
                    {
                        "heading": "Key Terms",
                        "heading_zh": "关键术语",
                        "items": [
                            {"en": "Systems analysis", "zh": "系统分析"},
                            {"en": "SDLC", "zh": "系统开发生命周期"},
                        ],
                    },
                    {
                        "heading": "Bilingual Sentences",
                        "heading_zh": "双语句型",
                        "items": [{
                            "en": "Systems analysis is the study of an existing system in order to improve it.",
                            "zh": "系统分析是研究现有系统，以便改进它。",
                        }],
                    },
                ],
                "sources": ["MIS413_Week1_Foundations_and_the_SDLC.pdf, pages 4-5"],
            }

            output = generate_report(course, report)

            self.assertTrue(output.exists())

    def test_generate_report_respects_chinese_first_language_mode(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            course = Path(temp_dir) / "course"
            course.mkdir()
            report = {
                "title": {"en": "Preview", "zh": "预习"},
                "course": "course",
                "language_mode": "zh-en",
                "summary": {"en": "English summary", "zh": "中文总结"},
                "sections": [],
            }
            output = generate_report(course, report)
            from docx import Document
            text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
            self.assertLess(text.index("中文总结"), text.index("English summary"))

    def test_generate_report_removes_xml_invalid_control_characters(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            course = Path(temp_dir) / "course"
            course.mkdir()
            report = {
                "title": "Control test",
                "course": "course",
                "summary": "Before\x0bAfter",
                "sections": [],
            }
            output = generate_report(course, report)
            from docx import Document
            self.assertEqual(Document(output).paragraphs[3].text, "BeforeAfter")
            self.assertEqual(output.suffix, ".docx")
            self.assertTrue(output.parent.name == "reports")
            self.assertGreater(output.stat().st_size, 0)


if __name__ == "__main__":
    unittest.main()
