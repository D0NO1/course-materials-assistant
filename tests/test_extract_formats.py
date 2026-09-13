import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parents[1] / "scripts"))

from extract_material import extract_material


class ExtractFormatTests(unittest.TestCase):
    def test_docx_preserves_paragraph_and_table_locators(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "assignment.docx"
            document = Document()
            document.add_heading("Assignment Requirements", level=1)
            document.add_paragraph("Submit a scope statement.")
            document.add_table(rows=1, cols=2).rows[0].cells[0].text = "Rubric"
            document.save(path)

            result = extract_material(path)

            self.assertEqual(result["status"], "extracted")
            self.assertTrue(any(item["type"] == "heading" for item in result["locators"]))
            self.assertTrue(any(item["type"] == "table" for item in result["locators"]))

    def test_pptx_preserves_slide_text_and_slide_locator(self):
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "lecture.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "SDLC"
            slide.placeholders[1].text = "Planning and analysis"
            presentation.save(path)

            result = extract_material(path)

            self.assertEqual(result["status"], "extracted")
            self.assertEqual(result["locators"][0]["type"], "slide")
            self.assertIn("Planning and analysis", result["content"]["slides"][0]["text"])

    def test_xlsx_preserves_sheet_headers_formulas_and_hidden_state(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "data.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Scores"
            sheet.append(["Student", "Score"])
            sheet.append(["A", 95])
            sheet["C1"] = "Total"
            sheet["C2"] = "=B2"
            hidden = workbook.create_sheet("Hidden")
            hidden.sheet_state = "hidden"
            workbook.save(path)

            result = extract_material(path)

            self.assertEqual(result["status"], "extracted")
            sheets = result["content"]["sheets"]
            self.assertEqual(sheets[0]["headers"][:2], ["Student", "Score"])
            self.assertEqual(sheets[0]["formulas"][0]["formula"], "=B2")
            self.assertTrue(sheets[1]["hidden"])

    def test_access_is_explicitly_inventory_only(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "database.accdb"
            path.write_bytes(b"placeholder")

            result = extract_material(path)

            self.assertEqual(result["status"], "inventory-only")
            self.assertTrue(any("macros" in warning for warning in result["warnings"]))


if __name__ == "__main__":
    unittest.main()
